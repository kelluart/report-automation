from pathlib import Path
import re
import pandas as pd
from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parents[1]
DATA_PATH = BASE_DIR / "data" / "input.csv"
OUTPUT_DIR = BASE_DIR / "output"
OUTPUT_DIR.mkdir(exist_ok=True)

EXCEL_PATH = OUTPUT_DIR / "report.xlsx"


# -------------------------
# Load
# -------------------------
def load_csv(path: Path) -> pd.DataFrame:
    sample = pd.read_csv(path, header=None, nrows=1)
    if sample.empty:
        return pd.DataFrame()

    first_val = sample.iat[0, 0]
    first_str = "" if pd.isna(first_val) else str(first_val).strip()
    date_headerless = re.match(r"^\d{4}-\d{2}-\d{2}$", first_str) is not None

    if date_headerless:
        df = pd.read_csv(path, header=None)
        col_names = ["date", "category", "amount", "description"]
        if df.shape[1] > len(col_names):
            extra = [f"col_{i}" for i in range(len(col_names) + 1, df.shape[1] + 1)]
            col_names.extend(extra)
        df.columns = col_names[: df.shape[1]]
        return df

    return pd.read_csv(path)


# -------------------------
# Clean (basic + predictable)
# -------------------------
def basic_clean(df: pd.DataFrame) -> tuple[pd.DataFrame, dict]:
    info = {}
    df = df.copy()

    # normalize headers
    df.columns = [c.strip().lower() for c in df.columns]

    # drop duplicates
    before = len(df)
    df = df.drop_duplicates()
    info["duplicates_removed"] = int(before - len(df))
    info["rows_before_dedupe"] = int(before)
    info["rows_after_dedupe"] = int(len(df))

    # best-effort conversions
    if "date" in df.columns:
        parsed = pd.to_datetime(df["date"], errors="coerce")
        if parsed.notna().any():
            df["date"] = parsed
        else:
            df["date"] = df["date"].astype("string")

    if "amount" in df.columns:
        parsed = pd.to_numeric(df["amount"], errors="coerce")
        if parsed.notna().any():
            df["amount"] = parsed
        else:
            df["amount"] = df["amount"].astype("string")

    return df, info


# -------------------------
# Organize Clean_Data (sorting)
# -------------------------
def organize_clean_data(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()

    sort_cols = []
    if "category" in df.columns:
        # ensure consistent ordering even with missing categories
        df["category"] = df["category"].astype("string")
        sort_cols.append("category")

    if "date" in df.columns:
        sort_cols.append("date")

    if "amount" in df.columns:
        sort_cols.append("amount")

    if sort_cols:
        df = df.sort_values(by=sort_cols, ascending=True, na_position="last")

    return df


# -------------------------
# Build Summary Tables
# -------------------------
def build_metrics_table(df: pd.DataFrame, info: dict) -> pd.DataFrame:
    null_matrix = df.isna() | df.eq("")
    rows = [
        ("rows_total", int(len(df))),
        ("columns_total", int(len(df.columns))),
        ("duplicates_removed", int(info.get("duplicates_removed", 0))),
        ("null_columns_total", int(null_matrix.all().sum())),
        ("null_cells_total", int(null_matrix.sum().sum())),
    ]
    return pd.DataFrame(rows, columns=["metric", "value"])


def _get_amount_numeric(df: pd.DataFrame) -> pd.Series | None:
    if "amount" not in df.columns:
        return None
    numeric = pd.to_numeric(df["amount"], errors="coerce")
    if numeric.notna().any():
        return numeric
    return None


def build_category_totals(df: pd.DataFrame) -> pd.DataFrame:
    if "category" not in df.columns:
        return pd.DataFrame(columns=["category", "count", "sum_amount"])

    categories = df["category"].fillna("NULL").replace("", "NULL").astype("string")
    amount_numeric = _get_amount_numeric(df)

    if amount_numeric is not None:
        temp = pd.DataFrame({"category": categories, "amount": amount_numeric})
        cat = (
            temp.groupby("category", dropna=False)
            .agg(count=("category", "count"), sum_amount=("amount", "sum"))
            .reset_index()
            .sort_values(by=["category"])
        )
    else:
        cat = (
            categories.to_frame("category")
            .groupby("category", dropna=False)
            .size()
            .reset_index(name="count")
            .sort_values(by=["category"])
        )
        cat["sum_amount"] = pd.NA

    return cat[["category", "count", "sum_amount"]]


def build_month_totals(df: pd.DataFrame) -> pd.DataFrame | None:
    if "date" not in df.columns:
        return None

    parsed = pd.to_datetime(df["date"], errors="coerce")
    if not parsed.notna().any():
        return None

    amount_numeric = _get_amount_numeric(df)
    month = parsed.dt.to_period("M").astype(str)
    temp = pd.DataFrame({"month": month})

    if amount_numeric is not None:
        temp["amount"] = amount_numeric
        m = (
            temp.dropna(subset=["month"])
            .groupby("month", dropna=False)
            .agg(count=("month", "count"), sum_amount=("amount", "sum"))
            .reset_index()
            .sort_values("month")
        )
    else:
        m = (
            temp.dropna(subset=["month"])
            .groupby("month", dropna=False)
            .size()
            .reset_index(name="count")
            .sort_values("month")
        )
        m["sum_amount"] = pd.NA

    return m[["month", "count", "sum_amount"]]


def build_summary_grid(
    metrics_df: pd.DataFrame,
    category_df: pd.DataFrame,
    month_df: pd.DataFrame | None,
) -> pd.DataFrame:
    rows: list[list[object]] = []

    rows.append(["metric", "value", ""])
    for _, row in metrics_df.iterrows():
        rows.append([row["metric"], row["value"], ""])

    rows.append(["", "", ""])

    rows.append(["Category Totals", "", ""])
    rows.append(["category", "count", "sum_amount"])
    for _, row in category_df.iterrows():
        rows.append([row["category"], row["count"], row["sum_amount"]])

    if month_df is not None:
        rows.append(["", "", ""])
        rows.append(["Month Totals", "", ""])
        rows.append(["month", "count", "sum_amount"])
        for _, row in month_df.iterrows():
            rows.append([row["month"], row["count"], row["sum_amount"]])

    return pd.DataFrame(rows)


# -------------------------
# Export (2 sheets only)
# -------------------------
def export_excel(clean_df: pd.DataFrame, summary_grid: pd.DataFrame,
                 path: Path) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        # Sheet 1: organized clean data
        clean_df.to_excel(writer, sheet_name="Clean_Data", index=False)

        # Sheet 2: Summary (stack tables with spacing)
        summary_grid.to_excel(
            writer, sheet_name="Summary", index=False, header=False, startrow=0
        )

        # Formatting (simple and readable)
        wb = writer.book

        def autofit(ws, max_width=45):
            for col_cells in ws.columns:
                max_len = 0
                col_letter = col_cells[0].column_letter
                for cell in col_cells:
                    val = "" if cell.value is None else str(cell.value)
                    if len(val) > max_len:
                        max_len = len(val)
                ws.column_dimensions[col_letter].width = min(max_len + 2, max_width)

        # Clean_Data formatting
        ws = wb["Clean_Data"]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        autofit(ws)

        # Date format if present
        headers = [cell.value for cell in ws[1]]
        if "date" in headers:
            date_col_idx = headers.index("date") + 1
            for r in range(2, ws.max_row + 1):
                cell = ws.cell(row=r, column=date_col_idx)
                if cell.value is not None:
                    cell.number_format = "yyyy-mm-dd"

        # Summary formatting
        ws2 = wb["Summary"]
        ws2.freeze_panes = "A2"
        autofit(ws2, max_width=55)


# -------------------------
# Main
# -------------------------
def main():
    print("SCRIPT STARTED")

    df_raw = load_csv(DATA_PATH)
    df_clean, info = basic_clean(df_raw)
    df_clean = organize_clean_data(df_clean)

    metrics_df = build_metrics_table(df_clean, info)
    category_df = build_category_totals(df_clean)
    month_df = build_month_totals(df_clean)
    summary_grid = build_summary_grid(metrics_df, category_df, month_df)

    print(f"Rows before dedupe: {info.get('rows_before_dedupe', len(df_raw))}")
    print(f"Rows after dedupe: {info.get('rows_after_dedupe', len(df_clean))}")
    print(f"Duplicates removed: {info.get('duplicates_removed', 0)}")

    export_excel(df_clean, summary_grid, EXCEL_PATH)

    print("DONE")
    print(f"Excel: {EXCEL_PATH}")

    wb = load_workbook(EXCEL_PATH)
    print(f"Sheetnames: {wb.sheetnames}")


if __name__ == "__main__":
    main()
